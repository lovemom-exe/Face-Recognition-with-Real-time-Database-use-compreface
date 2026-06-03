from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, time
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from .attendance_services import (
    ATTENDANCE_ABSENT,
    ATTENDANCE_EXCUSED,
    ATTENDANCE_LATE,
    ATTENDANCE_MANUAL,
    ATTENDANCE_ON_TIME,
    BusinessRuleError,
    METHOD_NONE,
)
from .models import AttendanceLog, AttendanceSession, ClassCourse, Student
from .serializers import dt


REPORT_SESSION_STATUSES = {"OPEN", "CLOSED", "LOCKED"}
COUNTED_PRESENT_STATUSES = {ATTENDANCE_ON_TIME, ATTENDANCE_LATE, ATTENDANCE_MANUAL}
RISK_ATTENDANCE_RATE = 80.0
RISK_ABSENT_COUNT = 3


def _round_rate(value: float) -> float:
    return round(value, 2)


def _rate(numerator: int, denominator: int) -> float:
    if denominator <= 0:
        return 0.0
    return _round_rate((numerator / denominator) * 100)


def _date_start(value: date | None) -> datetime | None:
    return datetime.combine(value, time.min) if value else None


def _date_end(value: date | None) -> datetime | None:
    return datetime.combine(value, time.max) if value else None


class ReportService:
    def __init__(self, db: Session) -> None:
        self.db = db

    def session_report(self, session_id: int) -> dict[str, Any]:
        session = self._get_session(session_id)
        students = self._students_for_class(session.class_course.class_id)
        logs_by_student = self._logs_by_student(session.id)
        student_rows = [self._session_student_row(session, student, logs_by_student.get(student.id)) for student in students]
        summary = self._summary_from_rows(student_rows)

        return {
            "session": self._session_payload(session),
            "class": self._class_payload(session.class_course),
            "course": self._course_payload(session.class_course),
            "summary": summary,
            "students": student_rows,
        }

    def student_report(
        self,
        student_id: int,
        *,
        course_id: int | None = None,
        class_id: int | None = None,
        from_date: date | None = None,
        to_date: date | None = None,
    ) -> dict[str, Any]:
        student = self.db.get(Student, student_id)
        if not student:
            raise BusinessRuleError("STUDENT_NOT_FOUND", "Không tìm thấy sinh viên.", status_code=404)

        effective_class_id = class_id or student.class_id
        if not effective_class_id:
            return {
                "student": self._student_payload(student),
                "summary": self._empty_student_summary(),
                "history": [],
            }

        sessions = self._sessions_for_class(
            effective_class_id,
            course_id=course_id,
            from_date=from_date,
            to_date=to_date,
        )
        logs = {
            log.session_id: log
            for log in self.db.query(AttendanceLog)
            .filter(AttendanceLog.student_id == student.id, AttendanceLog.session_id.in_([item.id for item in sessions] or [-1]))
            .all()
        }
        history = [self._student_history_row(session, logs.get(session.id)) for session in sessions]
        return {
            "student": self._student_payload(student),
            "summary": self._student_summary(history),
            "history": history,
        }

    def class_course_report(
        self,
        class_course_id: int,
        *,
        from_date: date | None = None,
        to_date: date | None = None,
    ) -> dict[str, Any]:
        class_course = self._get_class_course(class_course_id)
        students = self._students_for_class(class_course.class_id)
        sessions = self._sessions_for_class_course(class_course_id, from_date=from_date, to_date=to_date)

        session_rows = []
        for session in sessions:
            report = self.session_report(session.id)
            session_rows.append(
                {
                    "session_id": session.id,
                    "session_name": session.session_name,
                    "start_time": dt(session.start_time),
                    "status": session.status,
                    "total_students": report["summary"]["total_students"],
                    "on_time": report["summary"]["on_time"],
                    "late": report["summary"]["late"],
                    "absent": report["summary"]["absent"],
                    "excused": report["summary"]["excused"],
                    "manual": report["summary"]["manual"],
                    "not_recorded": report["summary"]["not_recorded"],
                    "attendance_rate": report["summary"]["attendance_rate"],
                }
            )

        student_rows = [self._class_course_student_row(student, sessions) for student in students]
        risk_students = [
            {
                **row,
                "warning": "Nghỉ nhiều" if row["absent"] >= RISK_ABSENT_COUNT else "Tỷ lệ chuyên cần thấp",
            }
            for row in student_rows
            if row["attendance_rate"] < RISK_ATTENDANCE_RATE or row["absent"] >= RISK_ABSENT_COUNT
        ]
        total_sessions = len(sessions)
        total_students = len(students)
        total_on_time = sum(row["on_time"] for row in student_rows)
        total_late = sum(row["late"] for row in student_rows)
        total_absent = sum(row["absent"] for row in student_rows)
        total_excused = sum(row["excused"] for row in student_rows)

        return {
            "class_course": self._class_course_payload(class_course),
            "summary": {
                "total_students": total_students,
                "total_sessions": total_sessions,
                "average_attendance_rate": _round_rate(sum(row["attendance_rate"] for row in student_rows) / total_students) if total_students else 0.0,
                "total_on_time": total_on_time,
                "total_late": total_late,
                "total_absent": total_absent,
                "total_excused": total_excused,
            },
            "sessions": session_rows,
            "students": student_rows,
            "risk_students": risk_students,
        }

    def export_session_excel(self, session_id: int) -> tuple[str, bytes]:
        report = self.session_report(session_id)
        workbook = Workbook()
        overview = workbook.active
        overview.title = "Tong quan"
        self._fill_key_value_sheet(
            overview,
            [
                ("Tên buổi", report["session"]["session_name"]),
                ("Lớp", report["class"]["class_name"]),
                ("Môn học", report["course"]["course_name"]),
                ("Thời gian bắt đầu", report["session"]["start_time"]),
                ("Trạng thái session", report["session"]["status"]),
                ("Tổng sinh viên", report["summary"]["total_students"]),
                ("Có mặt đúng giờ", report["summary"]["on_time"]),
                ("Đi muộn", report["summary"]["late"]),
                ("Vắng", report["summary"]["absent"]),
                ("Vắng có phép", report["summary"]["excused"]),
                ("Tỷ lệ chuyên cần", report["summary"]["attendance_rate"]),
            ],
        )

        detail = workbook.create_sheet("Danh sach diem danh")
        self._fill_table_sheet(
            detail,
            ["STT", "Mã sinh viên", "Họ tên", "Lớp", "Môn học", "Buổi", "Thời gian check-in", "Trạng thái", "Similarity", "Phương thức", "Ghi chú", "Recognition Event ID"],
            [
                [
                    index,
                    row["student_code"],
                    row["full_name"],
                    report["class"]["class_name"],
                    report["course"]["course_name"],
                    report["session"]["session_name"],
                    row["check_in_time"],
                    row["status"],
                    row["similarity"],
                    row["method"],
                    row["note"],
                    row["recognition_event_id"],
                ]
                for index, row in enumerate(report["students"], start=1)
            ],
        )
        self._format_workbook(workbook)
        return f"attendance_session_{session_id}.xlsx", self._workbook_bytes(workbook)

    def export_class_course_excel(
        self,
        class_course_id: int,
        *,
        from_date: date | None = None,
        to_date: date | None = None,
    ) -> tuple[str, bytes]:
        report = self.class_course_report(class_course_id, from_date=from_date, to_date=to_date)
        workbook = Workbook()
        overview = workbook.active
        overview.title = "Tong quan"
        self._fill_key_value_sheet(
            overview,
            [
                ("Lớp", report["class_course"]["class_name"]),
                ("Môn", report["class_course"]["course_name"]),
                ("Học kỳ", report["class_course"]["semester"]),
                ("Tổng sinh viên", report["summary"]["total_students"]),
                ("Tổng buổi", report["summary"]["total_sessions"]),
                ("Tỷ lệ chuyên cần trung bình", report["summary"]["average_attendance_rate"]),
                ("Tổng đúng giờ", report["summary"]["total_on_time"]),
                ("Tổng muộn", report["summary"]["total_late"]),
                ("Tổng vắng", report["summary"]["total_absent"]),
                ("Tổng có phép", report["summary"]["total_excused"]),
            ],
        )

        students_sheet = workbook.create_sheet("Theo sinh vien")
        self._fill_table_sheet(
            students_sheet,
            ["STT", "Mã sinh viên", "Họ tên", "Số buổi đúng giờ", "Số buổi muộn", "Số buổi vắng", "Số buổi có phép", "Tỷ lệ chuyên cần", "Cảnh báo"],
            [
                [
                    index,
                    row["student_code"],
                    row["full_name"],
                    row["on_time"],
                    row["late"],
                    row["absent"],
                    row["excused"],
                    row["attendance_rate"],
                    self._risk_warning(row),
                ]
                for index, row in enumerate(report["students"], start=1)
            ],
        )

        sessions_sheet = workbook.create_sheet("Theo buoi")
        self._fill_table_sheet(
            sessions_sheet,
            ["STT", "Tên buổi", "Thời gian", "Trạng thái", "Tổng sinh viên", "Đúng giờ", "Muộn", "Vắng", "Có phép", "Tỷ lệ chuyên cần"],
            [
                [
                    index,
                    row["session_name"],
                    row["start_time"],
                    row["status"],
                    row["total_students"],
                    row["on_time"],
                    row["late"],
                    row["absent"],
                    row["excused"],
                    row["attendance_rate"],
                ]
                for index, row in enumerate(report["sessions"], start=1)
            ],
        )
        self._format_workbook(workbook)
        return f"attendance_class_course_{class_course_id}.xlsx", self._workbook_bytes(workbook)

    def _get_session(self, session_id: int) -> AttendanceSession:
        session = self.db.get(AttendanceSession, session_id)
        if not session or not session.class_course:
            raise BusinessRuleError("SESSION_NOT_FOUND", "Không tìm thấy buổi điểm danh.", status_code=404)
        return session

    def _get_class_course(self, class_course_id: int) -> ClassCourse:
        class_course = self.db.get(ClassCourse, class_course_id)
        if not class_course:
            raise BusinessRuleError("CLASS_COURSE_NOT_FOUND", "Không tìm thấy lớp/môn.", status_code=404)
        return class_course

    def _students_for_class(self, class_id: int) -> list[Student]:
        return (
            self.db.query(Student)
            .filter(Student.class_id == class_id, Student.status == "ACTIVE")
            .order_by(Student.full_name)
            .all()
        )

    def _sessions_for_class(
        self,
        class_id: int,
        *,
        course_id: int | None,
        from_date: date | None,
        to_date: date | None,
    ) -> list[AttendanceSession]:
        query = (
            self.db.query(AttendanceSession)
            .join(AttendanceSession.class_course)
            .filter(ClassCourse.class_id == class_id, AttendanceSession.status.in_(REPORT_SESSION_STATUSES))
        )
        if course_id is not None:
            query = query.filter(ClassCourse.course_id == course_id)
        return self._apply_date_filter(query, from_date, to_date).order_by(AttendanceSession.start_time.desc()).all()

    def _sessions_for_class_course(
        self,
        class_course_id: int,
        *,
        from_date: date | None,
        to_date: date | None,
    ) -> list[AttendanceSession]:
        query = self.db.query(AttendanceSession).filter(
            AttendanceSession.class_course_id == class_course_id,
            AttendanceSession.status.in_(REPORT_SESSION_STATUSES),
        )
        return self._apply_date_filter(query, from_date, to_date).order_by(AttendanceSession.start_time.desc()).all()

    def _apply_date_filter(self, query: Any, from_date: date | None, to_date: date | None) -> Any:
        start = _date_start(from_date)
        end = _date_end(to_date)
        if start:
            query = query.filter(AttendanceSession.start_time >= start)
        if end:
            query = query.filter(AttendanceSession.start_time <= end)
        return query

    def _logs_by_student(self, session_id: int) -> dict[int, AttendanceLog]:
        return {
            log.student_id: log
            for log in self.db.query(AttendanceLog).filter(AttendanceLog.session_id == session_id).all()
        }

    def _session_student_row(self, session: AttendanceSession, student: Student, log: AttendanceLog | None) -> dict[str, Any]:
        if not log:
            status = ATTENDANCE_ABSENT if session.status in {"CLOSED", "LOCKED"} else "NOT_RECORDED"
            return {
                "student_id": student.id,
                "student_code": student.student_code,
                "full_name": student.full_name,
                "status": status,
                "check_in_time": None,
                "similarity": 0.0,
                "method": METHOD_NONE,
                "note": "",
                "recognition_event_id": None,
            }
        return {
            "student_id": student.id,
            "student_code": student.student_code,
            "full_name": student.full_name,
            "status": log.status,
            "check_in_time": dt(log.check_in_time),
            "similarity": log.similarity,
            "method": getattr(log, "method", None) or METHOD_NONE,
            "note": log.note,
            "recognition_event_id": log.recognition_event_id,
        }

    def _student_history_row(self, session: AttendanceSession, log: AttendanceLog | None) -> dict[str, Any]:
        class_course = session.class_course
        status = log.status if log else (ATTENDANCE_ABSENT if session.status in {"CLOSED", "LOCKED"} else "NOT_RECORDED")
        return {
            "session_id": session.id,
            "session_name": session.session_name,
            "course_name": class_course.course.course_name if class_course and class_course.course else None,
            "class_name": class_course.study_class.class_name if class_course and class_course.study_class else None,
            "start_time": dt(session.start_time),
            "status": status,
            "check_in_time": dt(log.check_in_time) if log else None,
            "similarity": log.similarity if log else 0.0,
            "method": getattr(log, "method", None) if log else METHOD_NONE,
            "note": log.note if log else "",
        }

    def _class_course_student_row(self, student: Student, sessions: list[AttendanceSession]) -> dict[str, Any]:
        if not sessions:
            return {
                "student_id": student.id,
                "student_code": student.student_code,
                "full_name": student.full_name,
                "on_time": 0,
                "late": 0,
                "absent": 0,
                "excused": 0,
                "manual": 0,
                "not_recorded": 0,
                "attendance_rate": 0.0,
            }
        logs = {
            log.session_id: log
            for log in self.db.query(AttendanceLog)
            .filter(AttendanceLog.student_id == student.id, AttendanceLog.session_id.in_([session.id for session in sessions]))
            .all()
        }
        counts = defaultdict(int)
        for session in sessions:
            log = logs.get(session.id)
            status = log.status if log else (ATTENDANCE_ABSENT if session.status in {"CLOSED", "LOCKED"} else "NOT_RECORDED")
            counts[status] += 1
        present = counts[ATTENDANCE_ON_TIME] + counts[ATTENDANCE_LATE] + counts[ATTENDANCE_MANUAL]
        return {
            "student_id": student.id,
            "student_code": student.student_code,
            "full_name": student.full_name,
            "on_time": counts[ATTENDANCE_ON_TIME],
            "late": counts[ATTENDANCE_LATE],
            "absent": counts[ATTENDANCE_ABSENT],
            "excused": counts[ATTENDANCE_EXCUSED],
            "manual": counts[ATTENDANCE_MANUAL],
            "not_recorded": counts["NOT_RECORDED"],
            "attendance_rate": _rate(present, len(sessions)),
        }

    def _summary_from_rows(self, rows: list[dict[str, Any]]) -> dict[str, Any]:
        total = len(rows)
        counts = defaultdict(int)
        for row in rows:
            counts[row["status"]] += 1
        present = counts[ATTENDANCE_ON_TIME] + counts[ATTENDANCE_LATE] + counts[ATTENDANCE_MANUAL]
        return {
            "total_students": total,
            "on_time": counts[ATTENDANCE_ON_TIME],
            "late": counts[ATTENDANCE_LATE],
            "absent": counts[ATTENDANCE_ABSENT],
            "excused": counts[ATTENDANCE_EXCUSED],
            "manual": counts[ATTENDANCE_MANUAL],
            "not_recorded": counts["NOT_RECORDED"],
            "attendance_rate": _rate(present, total),
            "absence_rate": _rate(counts[ATTENDANCE_ABSENT], total),
            "late_rate": _rate(counts[ATTENDANCE_LATE], total),
        }

    def _student_summary(self, history: list[dict[str, Any]]) -> dict[str, Any]:
        total = len(history)
        counts = defaultdict(int)
        for row in history:
            counts[row["status"]] += 1
        present = counts[ATTENDANCE_ON_TIME] + counts[ATTENDANCE_LATE] + counts[ATTENDANCE_MANUAL]
        return {
            "total_sessions": total,
            "on_time": counts[ATTENDANCE_ON_TIME],
            "late": counts[ATTENDANCE_LATE],
            "absent": counts[ATTENDANCE_ABSENT],
            "excused": counts[ATTENDANCE_EXCUSED],
            "manual": counts[ATTENDANCE_MANUAL],
            "not_recorded": counts["NOT_RECORDED"],
            "attendance_rate": _rate(present, total),
            "absence_rate": _rate(counts[ATTENDANCE_ABSENT], total),
            "late_rate": _rate(counts[ATTENDANCE_LATE], total),
        }

    def _empty_student_summary(self) -> dict[str, Any]:
        return {
            "total_sessions": 0,
            "on_time": 0,
            "late": 0,
            "absent": 0,
            "excused": 0,
            "manual": 0,
            "not_recorded": 0,
            "attendance_rate": 0.0,
            "absence_rate": 0.0,
            "late_rate": 0.0,
        }

    def _session_payload(self, session: AttendanceSession) -> dict[str, Any]:
        return {
            "id": session.id,
            "session_name": session.session_name,
            "status": session.status,
            "start_time": dt(session.start_time),
            "end_time": dt(session.end_time),
            "late_threshold_minutes": session.late_threshold_minutes,
        }

    def _class_payload(self, class_course: ClassCourse) -> dict[str, Any]:
        study_class = class_course.study_class
        return {
            "id": study_class.id if study_class else None,
            "class_code": study_class.class_code if study_class else None,
            "class_name": study_class.class_name if study_class else None,
        }

    def _course_payload(self, class_course: ClassCourse) -> dict[str, Any]:
        course = class_course.course
        return {
            "id": course.id if course else None,
            "course_code": course.course_code if course else None,
            "course_name": course.course_name if course else None,
        }

    def _class_course_payload(self, class_course: ClassCourse) -> dict[str, Any]:
        study_class = class_course.study_class
        course = class_course.course
        return {
            "id": class_course.id,
            "class_id": study_class.id if study_class else None,
            "class_code": study_class.class_code if study_class else None,
            "class_name": study_class.class_name if study_class else None,
            "course_id": course.id if course else None,
            "course_code": course.course_code if course else None,
            "course_name": course.course_name if course else None,
            "semester": class_course.semester,
        }

    def _student_payload(self, student: Student) -> dict[str, Any]:
        return {
            "id": student.id,
            "student_code": student.student_code,
            "full_name": student.full_name,
            "class_id": student.class_id,
            "class_name": student.study_class.class_name if student.study_class else None,
            "cohort": student.cohort,
            "major": student.major,
        }

    def _risk_warning(self, row: dict[str, Any]) -> str:
        if row["absent"] >= RISK_ABSENT_COUNT:
            return "Nghỉ nhiều"
        if row["attendance_rate"] < RISK_ATTENDANCE_RATE:
            return "Tỷ lệ chuyên cần thấp"
        return ""

    def _fill_key_value_sheet(self, sheet: Any, rows: list[tuple[str, Any]]) -> None:
        sheet.append(["Thông tin", "Giá trị"])
        for key, value in rows:
            sheet.append([key, value])

    def _fill_table_sheet(self, sheet: Any, headers: list[str], rows: list[list[Any]]) -> None:
        sheet.append(headers)
        for row in rows:
            sheet.append(row)

    def _format_workbook(self, workbook: Workbook) -> None:
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for column_cells in sheet.columns:
                max_length = 0
                column_letter = column_cells[0].column_letter
                for cell in column_cells:
                    max_length = max(max_length, len(str(cell.value)) if cell.value is not None else 0)
                sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 48)

    def _workbook_bytes(self, workbook: Workbook) -> bytes:
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
